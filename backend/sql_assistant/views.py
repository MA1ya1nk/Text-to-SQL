import csv
import hashlib
import io

from django.core.cache import cache
from django.db import connection
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rest_framework import status
from rest_framework.decorators import api_view
from rest_framework.response import Response

from .models import Favorite, QueryHistory
from .serializers import (
    ExecuteSQLRequestSerializer,
    ExplainRequestSerializer,
    FavoriteRequestSerializer,
    FavoriteSerializer,
    FollowUpRequestSerializer,
    QueryHistorySerializer,
    QueryRequestSerializer,
)
from .services.chart_detector import chart_detector
from .services.deterministic_sql_builder import deterministic_sql_builder
from .services.llm_client import LLMServiceError
from .services.nl_to_sql_service import nl_to_sql_service
from .services.query_executor import query_executor
from .services.query_explainer import query_explainer
from .services.schema_intent_guard import schema_intent_guard
from .services.schema_service import schema_service
from .services.sql_validator import sql_validator


def _generate_sql(question: str, context: str = "") -> str:
    dialect = "sqlite" if connection.vendor == "sqlite" else "postgresql"
    schema = schema_service.get_schema()
    if not schema_intent_guard.can_answer(question, schema):
        raise ValueError(
            "This question cannot be answered from the current schema. "
            "Please use available tables/fields or rephrase your question."
        )

    deterministic_sql = deterministic_sql_builder.build(question, dialect)
    if deterministic_sql:
        return sql_validator.validate_and_rewrite(deterministic_sql, dialect=dialect)

    matched_tables = schema_intent_guard.matched_tables(question, schema)
    schema_summary = schema_service.get_schema_summary(matched_tables=matched_tables)
    cache_basis = f"{dialect}|{question.strip().lower()}|{context.strip().lower()}|{schema_summary}"
    cache_key = "sql_generation_v1:" + hashlib.sha256(cache_basis.encode("utf-8")).hexdigest()
    cached_sql = cache.get(cache_key)
    if cached_sql:
        return cached_sql

    raw_sql = nl_to_sql_service.generate(
        question=question, schema_summary=schema_summary, dialect=dialect, context=context
    )
    validated_sql = sql_validator.validate_and_rewrite(raw_sql, dialect=dialect)
    cache.set(cache_key, validated_sql, 900)
    return validated_sql


def _execute_sql(question: str, sql: str):
    dialect = "sqlite" if connection.vendor == "sqlite" else "postgresql"
    safe_sql = sql_validator.validate_and_rewrite(sql, dialect=dialect)
    result = query_executor.execute(safe_sql)
    chart = chart_detector.detect(result["columns"], result["rows"])
    QueryHistory.objects.create(
        question=question,
        sql_query=safe_sql,
        explanation="",
        result_count=len(result["rows"]),
    )
    return {"sql": safe_sql, "chart": chart, **result}


def _error_response(exc: Exception) -> Response:
    if isinstance(exc, LLMServiceError):
        return Response(
            {"error": exc.user_message, "code": exc.code},
            status=exc.status_code,
        )
    if isinstance(exc, ValueError):
        return Response({"error": str(exc)}, status=status.HTTP_400_BAD_REQUEST)
    return Response(
        {"error": "Unable to process this request right now. Please try again."},
        status=status.HTTP_400_BAD_REQUEST,
    )


@api_view(["GET"])
def schema_view(request):
    return Response(schema_service.get_schema())


@api_view(["POST"])
def schema_refresh_view(request):
    return Response(schema_service.refresh_schema())


@api_view(["GET"])
def schema_suggestions_view(request):
    return Response({"suggestions": schema_service.get_suggestions()})


@api_view(["POST"])
def query_view(request):
    serializer = QueryRequestSerializer(data=request.data)
    serializer.is_valid(raise_exception=True)
    try:
        return Response(_execute_sql(serializer.validated_data["question"], _generate_sql(serializer.validated_data["question"])))
    except Exception as exc:
        return _error_response(exc)


@api_view(["POST"])
def followup_view(request):
    serializer = FollowUpRequestSerializer(data=request.data)
    serializer.is_valid(raise_exception=True)
    try:
        return Response(
            _execute_sql(
                serializer.validated_data["question"],
                _generate_sql(serializer.validated_data["question"], serializer.validated_data.get("context", "")),
            )
        )
    except Exception as exc:
        return _error_response(exc)


@api_view(["POST"])
def preview_view(request):
    serializer = FollowUpRequestSerializer(data=request.data)
    serializer.is_valid(raise_exception=True)
    try:
        sql = _generate_sql(serializer.validated_data["question"], serializer.validated_data.get("context", ""))
        return Response({"sql": sql})
    except Exception as exc:
        return _error_response(exc)


@api_view(["POST"])
def execute_view(request):
    serializer = ExecuteSQLRequestSerializer(data=request.data)
    serializer.is_valid(raise_exception=True)
    try:
        return Response(_execute_sql(serializer.validated_data["question"], serializer.validated_data["sql"]))
    except Exception as exc:
        return _error_response(exc)


@api_view(["POST"])
def explain_view(request):
    serializer = ExplainRequestSerializer(data=request.data)
    serializer.is_valid(raise_exception=True)
    try:
        explanation = query_explainer.explain(
            serializer.validated_data["sql"],
            question=serializer.validated_data.get("question", ""),
        )
        return Response({"explanation": explanation})
    except Exception as exc:
        return _error_response(exc)


@api_view(["GET"])
def history_view(request):
    data = QueryHistorySerializer(QueryHistory.objects.all()[:100], many=True).data
    return Response(data)


@api_view(["GET", "POST"])
def favorites_view(request):
    if request.method == "GET":
        data = FavoriteSerializer(Favorite.objects.all()[:100], many=True).data
        return Response(data)

    serializer = FavoriteRequestSerializer(data=request.data)
    serializer.is_valid(raise_exception=True)
    validated = serializer.validated_data
    favorite = Favorite.objects.create(
        title=validated.get("title") or validated["question"][:120],
        question=validated["question"],
        sql_query=validated["sql_query"],
    )
    return Response(FavoriteSerializer(favorite).data, status=status.HTTP_201_CREATED)


@api_view(["DELETE"])
def favorite_detail_view(request, favorite_id: int):
    try:
        favorite = Favorite.objects.get(pk=favorite_id)
    except Favorite.DoesNotExist:
        return Response({"error": "Favorite not found"}, status=status.HTTP_404_NOT_FOUND)
    favorite.delete()
    return Response(status=status.HTTP_204_NO_CONTENT)


@api_view(["POST"])
def export_view(request):
    columns = request.data.get("columns", [])
    rows = request.data.get("rows", [])
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(columns)
    writer.writerows(rows)
    response = HttpResponse(output.getvalue(), content_type="text/csv")
    response["Content-Disposition"] = 'attachment; filename="query_results.csv"'
    return response


@api_view(["POST"])
def export_excel_view(request):
    columns = request.data.get("columns", [])
    rows = request.data.get("rows", [])

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Query Results"
    header_fill = PatternFill(fill_type="solid", fgColor="EEF2FF")
    header_font = Font(bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    sanitized_columns = [str(col) if col is not None else "" for col in columns]
    worksheet.append(sanitized_columns)

    max_lengths = [len(str(col)) for col in sanitized_columns] or [0]

    for row in rows:
        normalized_row = list(row)[: len(sanitized_columns)]
        if len(normalized_row) < len(sanitized_columns):
            normalized_row.extend([None] * (len(sanitized_columns) - len(normalized_row)))

        sanitized_row = []
        for idx, value in enumerate(normalized_row):
            if value is None:
                cell_value = ""
            elif isinstance(value, (int, float, bool)):
                cell_value = value
            else:
                text = str(value).replace("\r\n", "\n").replace("\r", "\n")
                cell_value = text
            sanitized_row.append(cell_value)
            max_lengths[idx] = max(max_lengths[idx], len(str(cell_value)))
        worksheet.append(sanitized_row)

    if sanitized_columns:
        worksheet.auto_filter.ref = f"A1:{get_column_letter(len(sanitized_columns))}{worksheet.max_row}"
    worksheet.freeze_panes = "A2"

    for col_idx in range(1, len(sanitized_columns) + 1):
        column_letter = get_column_letter(col_idx)
        header_cell = worksheet[f"{column_letter}1"]
        header_cell.fill = header_fill
        header_cell.font = header_font
        header_cell.alignment = header_alignment

        width = min(max(max_lengths[col_idx - 1] + 2, 12), 60)
        worksheet.column_dimensions[column_letter].width = width

        for row_idx in range(2, worksheet.max_row + 1):
            worksheet[f"{column_letter}{row_idx}"].alignment = body_alignment

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="query_results.xlsx"'
    return response
